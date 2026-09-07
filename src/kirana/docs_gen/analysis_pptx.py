"""Business-analysis PowerPoint with real charts (python-pptx + matplotlib).

Slides: title → KPI summary → daily revenue trend → top items → payment mix →
GST slab breakdown → stock health & reorder suggestions → insights. Charts are
rendered as PNGs by matplotlib and embedded — real charts from real data, as
the task requires (no screenshots, no plain text).

Text formatting is applied at RUN level (not paragraph defaults) and word_wrap
is set explicitly, so the deck renders identically in PowerPoint, Keynote,
Google Slides and LibreOffice.
"""

from __future__ import annotations

import tempfile
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

from ..services import analytics_service, inventory_service, memory_service  # noqa: E402

IST = ZoneInfo("Asia/Kolkata")
OUT_DIR = Path("data/artifacts")

MAROON = RGBColor(0x5B, 0x23, 0x33)
TEAL = RGBColor(0x2A, 0x6F, 0x77)
GREY = RGBColor(0x55, 0x55, 0x55)
PALETTE = ["#5b2333", "#2a6f77", "#b3562e", "#7d5ba6", "#3b7a57", "#8c8c8c"]

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)


def _blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _textbox(slide, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    return tf


def _line(tf, text, size, *, bold=False, color=GREY, first=False, space_before=0):
    """Append a paragraph with run-level formatting (portable across renderers)."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = PP_ALIGN.LEFT
    if space_before:
        p.space_before = Pt(space_before)
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return p


def _title(slide, text: str, sub: str | None = None):
    tf = _textbox(slide, Inches(0.6), Inches(0.35), Inches(12.1), Inches(1.0))
    _line(tf, text, 28, bold=True, color=MAROON, first=True)
    if sub:
        _line(tf, sub, 13, color=GREY)


def _kpi(slide, x, label, value, color=MAROON):
    tf = _textbox(slide, x, Inches(1.9), Inches(3.0), Inches(1.5))
    _line(tf, value, 34, bold=True, color=color, first=True)
    _line(tf, label, 13, color=GREY)


def _chart_png(fig) -> str:
    tmp = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
    fig.savefig(tmp.name, dpi=160, bbox_inches="tight")
    plt.close(fig)
    return tmp.name


def generate_analysis_deck(date_from: str, date_to: str) -> str:
    data = analytics_service.sales_range(date_from, date_to)
    low = inventory_service.low_stock()
    reorder = analytics_service.reorder_suggestions()
    prefs = memory_service.get_preferences()
    shop = prefs.get("shop_name", "Kirana Store")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUT_DIR / f"analysis_{date_from}_to_{date_to}.pptx"

    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    t = data["totals"]
    khata_total = sum(b.get("amount", 0) or 0
                      for b in data["by_payment_mode"] if b["payment_mode"] == "khata")

    # ---- 1. title -----------------------------------------------------------
    s = _blank_slide(prs)
    tf = _textbox(s, Inches(0.8), Inches(2.5), Inches(11.7), Inches(2.4))
    _line(tf, f"{shop} — Sales Analysis", 40, bold=True, color=MAROON, first=True)
    _line(tf, f"{date_from} to {date_to}", 18, color=GREY, space_before=8)
    _line(tf, f"Generated {datetime.now(IST).strftime('%d %b %Y, %I:%M %p IST')} "
              f"by the Supermarket Ops Agent", 12, color=TEAL, space_before=6)

    # ---- 2. KPI summary -----------------------------------------------------
    s = _blank_slide(prs)
    _title(s, "Period at a Glance")
    _kpi(s, Inches(0.6), "Revenue", f"₹{t['revenue']:,.0f}")
    _kpi(s, Inches(3.9), "Bills cut", f"{t['bills']}")
    _kpi(s, Inches(7.2), "GST collected", f"₹{t['tax']:,.0f}", TEAL)
    _kpi(s, Inches(10.5), "Sold on khata", f"₹{khata_total:,.0f}", TEAL)
    avg = t["revenue"] / t["bills"] if t["bills"] else 0
    tf = _textbox(s, Inches(0.6), Inches(3.9), Inches(12.1), Inches(1.6))
    _line(tf, f"Average bill value: ₹{avg:,.0f}", 16, color=GREY, first=True)
    if low:
        _line(tf, f"{len(low)} item(s) at or below reorder level.", 16, color=GREY,
              space_before=8)

    # ---- 3. daily revenue trend --------------------------------------------
    if data["daily"]:
        fig, ax = plt.subplots(figsize=(10, 4.2))
        ax.bar([r["d"][5:] for r in data["daily"]],
               [r["revenue"] or 0 for r in data["daily"]], color=PALETTE[0])
        ax.set_ylabel("Revenue (₹)")
        ax.set_title("Daily revenue")
        ax.spines[["top", "right"]].set_visible(False)
        s = _blank_slide(prs)
        _title(s, "Daily Revenue Trend")
        s.shapes.add_picture(_chart_png(fig), Inches(0.9), Inches(1.6), width=Inches(11.5))

    # ---- 4. top items -------------------------------------------------------
    if data["top_items"]:
        items = data["top_items"][:8][::-1]
        fig, ax = plt.subplots(figsize=(10, 4.5))
        ax.barh([i["name"] for i in items], [i["revenue"] for i in items], color=PALETTE[1])
        ax.set_xlabel("Revenue (₹)")
        ax.set_title("Top items by revenue")
        ax.spines[["top", "right"]].set_visible(False)
        s = _blank_slide(prs)
        _title(s, "Top Sellers")
        s.shapes.add_picture(_chart_png(fig), Inches(0.9), Inches(1.5), width=Inches(11.5))

    # ---- 5. payment mix -----------------------------------------------------
    if data["by_payment_mode"]:
        fig, ax = plt.subplots(figsize=(5.5, 4.5))
        ax.pie([m["amount"] or 0 for m in data["by_payment_mode"]],
               labels=[(m["payment_mode"] or "?").upper() for m in data["by_payment_mode"]],
               autopct="%1.0f%%", colors=PALETTE, textprops={"fontsize": 11})
        ax.set_title("Collections by payment mode")
        s = _blank_slide(prs)
        _title(s, "Payment Mix", "Cash vs UPI vs Card vs Khata")
        s.shapes.add_picture(_chart_png(fig), Inches(4.2), Inches(1.6), height=Inches(5.2))

    # ---- 6. GST slabs -------------------------------------------------------
    if data["by_gst_slab"]:
        fig, ax = plt.subplots(figsize=(8, 4))
        ax.bar([f"{r['gst_rate']:g}%" for r in data["by_gst_slab"]],
               [r["tax"] or 0 for r in data["by_gst_slab"]], color=PALETTE[2])
        ax.set_ylabel("Tax collected (₹)")
        ax.set_title("GST collected by slab (CGST + SGST)")
        ax.spines[["top", "right"]].set_visible(False)
        s = _blank_slide(prs)
        _title(s, "GST Breakdown", "Per-slab tax liability for the period")
        s.shapes.add_picture(_chart_png(fig), Inches(2.4), Inches(1.7), width=Inches(8.6))

    # ---- 7. stock health ----------------------------------------------------
    s = _blank_slide(prs)
    _title(s, "Stock Health & Reorder", "Suggestions from sales velocity vs stock on hand")
    tf = _textbox(s, Inches(0.6), Inches(1.7), Inches(12.1), Inches(5.2))
    if reorder:
        for i, r in enumerate(reorder[:9]):
            left = (f"{r['days_of_stock_left']:g} days of stock left"
                    if r["days_of_stock_left"] is not None else "below reorder level")
            _line(tf, f"•  {r['product']} — {r['in_stock']:g} {r['unit']} on hand, "
                      f"{left}  →  order ~{r['suggested_order_qty']:g}",
                  15, color=GREY, first=(i == 0), space_before=4)
    else:
        _line(tf, "All items comfortably stocked — nothing urgent to reorder.",
              16, color=GREY, first=True)
    if low:
        _line(tf, f"At or below reorder level: {', '.join(x['name'] for x in low[:8])}",
              13, color=GREY, space_before=14)

    # ---- 8. insights --------------------------------------------------------
    insights = []
    if data["top_items"]:
        top = data["top_items"][0]
        insights.append(f"{top['name']} led the period with ₹{top['revenue']:,.0f} in sales.")
    if data["by_payment_mode"]:
        biggest = max(data["by_payment_mode"], key=lambda m: m["amount"] or 0)
        share = (biggest["amount"] or 0) / t["revenue"] * 100 if t["revenue"] else 0
        insights.append(f"{(biggest['payment_mode'] or '?').upper()} was the dominant payment "
                        f"mode at {share:.0f}% of collections.")
    if khata_total:
        insights.append(f"₹{khata_total:,.0f} of sales went on khata — "
                        f"worth chasing collections.")
    if reorder:
        insights.append(f"{len(reorder)} item(s) need reordering soon; fastest-moving is "
                        f"{reorder[0]['product']}.")
    if not insights:
        insights.append("No finalized sales in this period yet.")

    s = _blank_slide(prs)
    _title(s, "Insights")
    tf = _textbox(s, Inches(0.6), Inches(1.7), Inches(12.1), Inches(5.2))
    for i, text in enumerate(insights):
        _line(tf, f"•  {text}", 17, color=GREY, first=(i == 0), space_before=8)

    prs.save(str(path))
    return str(path)
